from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from loguru import logger

from .arxiv_email_parser import EmailPaper
from .paper_analyzer import PaperAnalysis

_SUB_SCORE_FIELDS = [
    ("research_problem_score", "当前研究问题相关度"),
    ("method_transfer_score", "方法可迁移性"),
    ("acoustic_modality_score", "声学适配程度"),
    ("novelty_score", "可验证的改进启发"),
    ("experiment_feasibility_score", "实验可行性与协议兼容性"),
]


def write_report(
    papers: list[EmailPaper],
    analyses: list[PaperAnalysis],
    output_dir: str | Path,
    report_date: date | None = None,
    diagnostics: list[str] | None = None,
    llm_diagnostics: dict | None = None,
) -> Path:
    report_date = report_date or date.today()
    output_path = Path(output_dir) / f"{report_date:%Y-%m-%d}_paper_triage_report.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 高分论文排在最前面,后续所有板块(快速筛选表、详细分析、清单)都复用这个顺序,
    # 保证"打开报告先看到最相关的论文"。
    papers, analyses = _sort_by_relevance(papers, analyses)

    document = Document()
    # python-docx 默认模板的 zoom 缺少 OOXML 必需的 percent 属性。
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    _set_default_font(document)
    _set_landscape_layout(document)
    document.add_heading("每日 arXiv 论文迁移可行性分析报告", level=0)

    _add_overview(document, papers, analyses, report_date, diagnostics or [], llm_diagnostics or {})
    _add_quick_table(document, papers, analyses)
    _add_inspiration_list(document, papers, analyses)
    _add_details(document, papers, analyses)
    _add_reading_list(document, "今日精读清单", papers, analyses, "精读")
    _add_reading_list(document, "今日略读清单", papers, analyses, "略读")
    _add_reading_list(document, "今日待核验清单", papers, analyses, "待核验")
    _add_reading_list(document, "今日跳过清单", papers, analyses, "跳过")
    _add_tags(document, analyses)
    _add_experiment_ideas(document, papers, analyses)

    document.save(output_path)
    logger.info("Wrote paper triage report to {}", output_path)
    return output_path


def _sort_by_relevance(
    papers: list[EmailPaper], analyses: list[PaperAnalysis]
) -> tuple[list[EmailPaper], list[PaperAnalysis]]:
    if not papers:
        return papers, analyses
    paired = sorted(zip(papers, analyses), key=lambda item: item[1].relevance_score, reverse=True)
    sorted_papers, sorted_analyses = zip(*paired)
    return list(sorted_papers), list(sorted_analyses)


def _set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)


def _set_landscape_layout(document: Document) -> None:
    for section in document.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)


def _add_overview(
    document: Document,
    papers: list[EmailPaper],
    analyses: list[PaperAnalysis],
    report_date: date,
    diagnostics: list[str],
    llm_diagnostics: dict,
) -> None:
    document.add_heading("今日总览", level=1)
    document.add_paragraph(f"报告日期：{report_date:%Y-%m-%d}")
    document.add_paragraph(f"解析论文数量：{len(papers)}")
    counts = Counter(a.reading_recommendation for a in analyses)
    document.add_paragraph(
        f"阅读建议统计：精读 {counts.get('精读', 0)} 篇，"
        f"略读 {counts.get('略读', 0)} 篇，待核验 {counts.get('待核验', 0)} 篇，跳过 {counts.get('跳过', 0)} 篇。"
    )
    if analyses:
        valid = [a for a in analyses if not a.error]
        if valid:
            avg_score = sum(a.relevance_score for a in valid) / len(valid)
            document.add_paragraph(f"有效分析平均综合分：{avg_score:.1f}/10（排除分析失败；摘要评分为初评）")
        else:
            document.add_paragraph("有效分析平均综合分：不可用（全部分析失败）")
    document.add_paragraph(
        "评分说明（GZSL 组件价值评分 v2）：综合分由 当前研究问题相关度、方法可迁移性、声学适配程度、可验证的改进启发、"
        "实验可行性与协议兼容性 五项分项分按固定权重加权计算，权重见 paper_triage_config.yaml 的 "
        "scoring_weights；风险分单独展示，不计入综合分。每篇论文的具体分项分见下方逐篇详细分析。"
    )

    document.add_paragraph("阅读建议：精读需综合分≥7.5、组件明确、证据充分且协议兼容；略读通常需≥5.5及明确组件，强相关局部组件可优先略读。证据不足或协议未知标待核验，协议不兼容标跳过。所有分析基于提供的邮件摘要/TLDR，未读取全文。")

    failures = [analysis for analysis in analyses if analysis.error]
    if failures:
        by_category = Counter(analysis.error_category or "未知错误" for analysis in failures)
        failure_summary = "；".join(f"{name} {count} 篇" for name, count in by_category.items())
        document.add_paragraph(f"LLM 分析失败统计：{failure_summary}")

    if llm_diagnostics:
        document.add_paragraph("LLM 配置摘要：")
        document.add_paragraph(
            "；".join(
                [
                    f"provider={llm_diagnostics.get('provider', 'unknown')}",
                    f"base_url={llm_diagnostics.get('base_url', 'unknown')}",
                    f"model={llm_diagnostics.get('model', 'unknown')}",
                    f"key_source={llm_diagnostics.get('key_source', 'unknown')}",
                ]
            ),
            style="List Bullet",
        )
        for warning in llm_diagnostics.get("warnings", []) or []:
            document.add_paragraph(f"警告：{warning}", style="List Bullet")

    if diagnostics:
        document.add_paragraph("诊断信息：")
        for item in diagnostics:
            document.add_paragraph(item, style="List Bullet")


def _add_quick_table(document: Document, papers: list[EmailPaper], analyses: list[PaperAnalysis]) -> None:
    document.add_heading("快速筛选表", level=1)
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(8.0), Cm(1.6), Cm(3.6), Cm(2.2), Cm(2.0), Cm(1.6), Cm(3.8)]
    headers = ["标题", "综合分", "方向", "迁移可行性", "阅读建议", "风险", "标签"]
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True)
        table.rows[0].cells[index].width = widths[index]

    if not papers:
        row = table.add_row().cells
        row[0].text = "未解析到论文"
        return

    for paper, analysis in zip(papers, analyses):
        row = table.add_row().cells
        risk_text = f"{analysis.risk_score:g}" if analysis.risk_score >= 7 else "—"
        values = [
            _shorten(paper.title, 70),
            "未评分" if analysis.error else f"{analysis.relevance_score:g}",
            _shorten(analysis.matched_research_direction, 32),
            analysis.transfer_feasibility,
            f"{analysis.reading_recommendation}（证据{analysis.evidence_sufficiency}）",
            risk_text,
            ", ".join(analysis.keyword_tags),
        ]
        for index, value in enumerate(values):
            cell = row[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if index == 1:
                _set_cell_text(cell, value, bold=True, color=_score_color(analysis.relevance_score))
            elif index == 3:
                _set_cell_text(cell, value, bold=True, color=_feasibility_color(value))
            elif index == 4:
                _set_cell_text(cell, value, bold=True, color=_recommendation_color(analysis.reading_recommendation))
            elif index == 5 and risk_text != "—":
                _set_cell_text(cell, value, bold=True, color=RGBColor(0xB0, 0x00, 0x20))
            else:
                _set_cell_text(cell, value)


def _add_inspiration_list(document: Document, papers: list[EmailPaper], analyses: list[PaperAnalysis]) -> None:
    """专门回答"今天这些论文里有没有能借鉴到我 GZSL 课题里的新点子"这个问题。"""
    document.add_heading("今日灵感清单", level=1)
    document.add_paragraph("按综合分从高到低，只列出 LLM 认为有具体新颖点的论文（排除「无明显新颖点」）。")
    selected = [
        (paper, analysis)
        for paper, analysis in zip(papers, analyses)
        if analysis.concrete_component and not analysis.error
        and analysis.inspiration_note and not analysis.inspiration_note.startswith(("无明显", "待人工复核"))
    ]
    if not selected:
        document.add_paragraph("今日没有被判定为有明显新颖点的论文。")
        return
    for paper, analysis in selected:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(f"{paper.title}（综合分 {analysis.relevance_score:g}/10）：")
        run.bold = True
        paragraph.add_run(analysis.inspiration_note)


def _add_details(document: Document, papers: list[EmailPaper], analyses: list[PaperAnalysis]) -> None:
    document.add_heading("逐篇论文详细分析", level=1)
    if not papers:
        document.add_paragraph("未解析到可分析论文。")
        return

    for index, (paper, analysis) in enumerate(zip(papers, analyses), start=1):
        document.add_heading(f"{index}. {paper.title}", level=2)
        link_paragraph = document.add_paragraph("arXiv 链接：")
        _add_hyperlink(link_paragraph, paper.arxiv_url, paper.arxiv_url)
        score_text = "未评分" if analysis.error else f"{analysis.relevance_score:g}/10（摘要初评）"
        document.add_paragraph(f"综合分：{score_text}（{analysis.reading_recommendation}）")
        document.add_paragraph(f"作用环节：{analysis.research_stage}")
        document.add_paragraph(f"阅读依据：{analysis.evidence_source}")
        document.add_paragraph(f"证据充分度：{analysis.evidence_sufficiency}；证据缺口：{analysis.evidence_gaps}")
        document.add_paragraph(f"监督要求：{analysis.supervision_requirements}")
        document.add_paragraph(f"协议兼容性：{analysis.protocol_compatibility}；{analysis.protocol_reason}")
        document.add_paragraph(f"具体可迁移组件：{'已识别' if analysis.concrete_component else '尚未明确'}")
        document.add_paragraph(f"分项评分：{_format_sub_scores(analysis)}")
        document.add_paragraph(f"评分依据：{analysis.score_rationale}")
        document.add_paragraph(f"评分计算过程：{analysis.decision_reason}")
        document.add_paragraph(f"最匹配研究方向：{analysis.matched_research_direction}")
        document.add_paragraph(f"核心贡献：{analysis.core_contribution}")
        document.add_paragraph(f"方法类型：{analysis.method_type}")
        document.add_paragraph(f"技术要素：{_format_flags(analysis.technique_flags)}")
        document.add_paragraph(
            f"是否可迁移到水下目标识别：{'是' if analysis.transferable_to_underwater else '否'}"
            f"（迁移可行性：{analysis.transfer_feasibility}）"
        )
        document.add_paragraph(f"可以迁移的部分：{analysis.transferable_parts}")
        document.add_paragraph(f"不容易迁移的部分：{analysis.hard_to_transfer_parts}")
        document.add_paragraph(f"DeepShip / ShipsEar 最小实验方案：{analysis.minimal_experiment}")
        inspiration_paragraph = document.add_paragraph()
        inspiration_run = inspiration_paragraph.add_run("给你的灵感：")
        inspiration_run.bold = True
        inspiration_paragraph.add_run(analysis.inspiration_note)
        if analysis.risk_score >= 7:
            risk_paragraph = document.add_paragraph()
            risk_run = risk_paragraph.add_run(f"风险提示：risk_score {analysis.risk_score:g}/10 偏高，注意复现/迁移成本。")
            risk_run.bold = True
            risk_run.font.color.rgb = RGBColor(0xB0, 0x00, 0x20)
        document.add_paragraph(f"关键词标签：{', '.join(analysis.keyword_tags)}")
        if analysis.error:
            document.add_paragraph(f"自动分析错误类型：{analysis.error_category or '未知错误'}")
            document.add_paragraph(f"自动分析错误详情：{analysis.error}")


def _add_reading_list(
    document: Document,
    heading: str,
    papers: list[EmailPaper],
    analyses: list[PaperAnalysis],
    recommendation: str,
) -> None:
    document.add_heading(heading, level=1)
    selected = [
        (paper, analysis)
        for paper, analysis in zip(papers, analyses)
        if analysis.reading_recommendation == recommendation
    ]
    if not selected:
        document.add_paragraph("无")
        return
    for paper, analysis in selected:
        score_text = "未评分" if analysis.error else f"{analysis.relevance_score:g}/10"
        document.add_paragraph(
            f"{paper.title}（{score_text}，{analysis.transfer_feasibility}）",
            style="List Bullet",
        )


def _add_tags(document: Document, analyses: list[PaperAnalysis]) -> None:
    document.add_heading("可加入 Zotero 的标签建议", level=1)
    tags = []
    for analysis in analyses:
        tags.extend(analysis.keyword_tags)
    unique_tags = sorted(set(tag for tag in tags if tag))
    document.add_paragraph(", ".join(unique_tags) if unique_tags else "无")


def _add_experiment_ideas(document: Document, papers: list[EmailPaper], analyses: list[PaperAnalysis]) -> None:
    document.add_heading("后续实验启发", level=1)
    selected = [
        (paper, analysis)
        for paper, analysis in zip(papers, analyses)
        if analysis.reading_recommendation in {"精读", "略读"}
    ]
    if not selected:
        document.add_paragraph("暂无明确实验启发。")
        return
    for paper, analysis in selected:
        document.add_paragraph(f"{paper.title}：{analysis.minimal_experiment}", style="List Bullet")


def _format_sub_scores(analysis: PaperAnalysis) -> str:
    parts = [f"{label} {getattr(analysis, field):g}" for field, label in _SUB_SCORE_FIELDS]
    parts.append(f"风险分 {analysis.risk_score:g}")
    return "，".join(parts)


def _set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text or "")
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shorten(text: str, limit: int) -> str:
    clean = " ".join((text or "").split())
    if len(clean) <= limit:
        return clean
    return clean[: limit - 1].rstrip() + "…"


def _score_color(value: float) -> RGBColor:
    if value >= 7.5:
        return RGBColor(0x1B, 0x7F, 0x3A)
    if value >= 5.5:
        return RGBColor(0x9A, 0x67, 0x00)
    return RGBColor(0x66, 0x66, 0x66)


def _feasibility_color(value: str) -> RGBColor:
    if value == "高":
        return RGBColor(0x1B, 0x7F, 0x3A)
    if value == "中":
        return RGBColor(0x9A, 0x67, 0x00)
    return RGBColor(0xB0, 0x00, 0x20)


def _recommendation_color(value: str) -> RGBColor:
    if value == "精读":
        return RGBColor(0x1B, 0x7F, 0x3A)
    if value == "略读":
        return RGBColor(0x9A, 0x67, 0x00)
    return RGBColor(0x66, 0x66, 0x66)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _format_flags(flags: dict[str, bool]) -> str:
    positives = [name for name, enabled in flags.items() if enabled]
    return ", ".join(positives) if positives else "未明显涉及"
