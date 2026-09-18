from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT

from paper_triage.arxiv_email_parser import EmailPaper
from paper_triage.paper_analyzer import PaperAnalysis
from paper_triage.report_writer import write_report


def test_write_report_creates_docx_with_required_sections(tmp_path):
    paper = EmailPaper(
        title=(
            "Open-set prototype learning for acoustic classification with cross-domain "
            "few-shot adaptation and robust unknown rejection"
        ),
        abstract="A short abstract.",
        arxiv_url="https://arxiv.org/abs/2401.00001",
    )
    analysis = PaperAnalysis(
        relevance_score=8,
        matched_research_direction="开放集识别 OSR",
        core_contribution="提出开放集原型学习方法。",
        method_type="prototype learning",
        technique_flags={"open-set": True, "prototype learning": True},
        transferable_to_underwater=True,
        transfer_feasibility="高",
        transferable_parts="原型分类和拒识阈值。",
        hard_to_transfer_parts="依赖图像增强的部分。",
        minimal_experiment="在 DeepShip 上提取 Log-Mel 后测试。",
        reading_recommendation="精读",
        keyword_tags=["OSR", "prototype", "DeepShip"],
        error_category="JSON 解析失败",
        error="JSON 解析失败: response preview",
    )

    path = write_report(
        [paper],
        [analysis],
        tmp_path,
        report_date=date(2026, 6, 4),
        llm_diagnostics={
            "provider": "SiliconFlow",
            "base_url": "https://api.siliconflow.cn/v1",
            "model": "Qwen/Qwen2.5-7B-Instruct",
            "key_source": "SILICONFLOW_API_KEY",
            "warnings": ["Set PAPER_TRIAGE_MODEL"],
        },
    )

    assert path.exists()
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert "每日 arXiv 论文迁移可行性分析报告" in text
    assert "今日总览" in text
    assert "逐篇论文详细分析" in text
    assert "今日精读清单" in text
    assert "provider=SiliconFlow" in text
    assert "自动分析错误类型：JSON 解析失败" in text
    assert document.tables[0].cell(1, 0).text.endswith("…")


def test_report_includes_evidence_and_pending_list(tmp_path):
    paper = EmailPaper("OOD gating", "TLDR", "https://arxiv.org/abs/1")
    analysis = PaperAnalysis(7.8, "GZSL 门控", "动态层选择", "OOD",
                             research_problem_score=9, reading_recommendation="待核验",
                             research_stage="已见/未见门控", supervision_requirements="外部 OOD 数据未知",
                             evidence_gaps="训练监督待核验", inspiration_note="无明显新颖点；只有类比")
    path = write_report([paper], [analysis], tmp_path)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for expected in ["当前研究问题相关度", "待核验 1 篇", "今日待核验清单", "作用环节：已见/未见门控",
                     "阅读依据：邮件摘要/TLDR", "监督要求：外部 OOD 数据未知", "协议兼容性：待核验"]:
        assert expected in text
    assert "今日没有被判定为有明显新颖点的论文" in text
    assert "证据低" in doc.tables[0].cell(1, 4).text
